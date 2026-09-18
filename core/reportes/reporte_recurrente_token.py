# -*- coding: utf-8 -*-
"""
core/reportes/reporte_recurrente_token.py

Reporte recurrente de facturación electrónica (DIAN) por empresa y período.
Genera 3 hojas con el diseño estándar:
    1. Resumen IVA   (ventas netas, compras netas, IVA a pagar)
    2. Ventas diarias (por fecha, con día de la semana y devoluciones)
    3. Compras + IVA por proveedor

Uso:
    from core.reportes.reporte_recurrente_token import generar_reporte
    generar_reporte(
        token_path="TOKEN.xlsx",
        nit_empresa="900451388",
        anio=2026, mes=9,
        dia_ini=1, dia_fin=17,               # opcional; None = mes completo
        titulo_empresa="SILLA TRES S.A.S",
        nit_formateado="900.451.388",
        salida="Reporte_SILLA3_Sep_1-17.xlsx",
    )
"""
import io
import datetime
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Índices de columna del token DIAN (formato estándar de 32 columnas) ---
I_TIPO, I_FE, I_EMI, I_NEMI, I_REC, I_IVA, I_TOT = 0, 7, 9, 10, 11, 13, 29
OTROS = list(range(14, 26))   # ICA, IC, INC, IBUA, ICUI, etc. -> "otros impuestos"
DIAS = ['LUNES', 'MARTES', 'MIÉRCOLES', 'JUEVES', 'VIERNES', 'SÁBADO', 'DOMINGO']

# --- Paleta / estilos ---
AR = "Arial"
NAVY, BLUE, LBLUE = "1F3864", "2E5496", "D9E1F2"
GOLD, GREEN, GREENF, WHITE = "FFF2CC", "C6EFCE", "006100", "FFFFFF"
_thin = Side(style="thin", color="BFBFBF")
_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _num(v):
    try:
        return float(str(v).replace(',', ''))
    except Exception:
        return 0.0


def _parse_fecha(fe):
    s = str(fe)[:10]
    for fmt in ('%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.datetime.strptime(s, fmt)
        except Exception:
            pass
    return None


def _es_nota_credito(tipo):
    t = str(tipo).lower()
    return 'cr' in t and 'dito' in t


def _S(c, v, b=False, sz=10, col="000000", fill=None, al="left", num=False, wrap=False):
    c.value = v
    c.font = Font(name=AR, bold=b, size=sz, color=col)
    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.border = _BD
    if num:
        c.number_format = '#,##0'


def procesar_token(token_path, nit_empresa, anio, mes, dia_ini=None, dia_fin=None):
    """Lee el token y devuelve (ventas_fact, ventas_nc, proveedores, resumen) por período."""
    wb = load_workbook(token_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    CO = str(nit_empresa).strip()
    vf = defaultdict(lambda: {'n': 0, 'base': 0.0, 'iva': 0.0, 'otros': 0.0, 'tot': 0.0})
    vn = defaultdict(lambda: {'n': 0, 'base': 0.0, 'iva': 0.0, 'otros': 0.0, 'tot': 0.0})
    prov = defaultdict(lambda: {'nom': '', 'nf': 0, 'nc': 0, 'base': 0.0, 'iva': 0.0, 'otros': 0.0, 'tot': 0.0})
    res = {k: {'n': 0, 'base': 0, 'iva': 0, 'otros': 0, 'tot': 0} for k in ('vf', 'vn', 'cf', 'cn')}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[I_TIPO] in (None, ''):
            continue
        tipo = str(r[I_TIPO])
        if tipo.startswith('Application'):
            continue
        d = _parse_fecha(r[I_FE])
        if not d or d.year != anio or d.month != mes:
            continue
        if dia_ini and d.day < dia_ini:
            continue
        if dia_fin and d.day > dia_fin:
            continue
        es_nc = _es_nota_credito(tipo)
        emi = str(r[I_EMI] or '').strip()
        rec = str(r[I_REC] or '').strip()
        iva = _num(r[I_IVA]); tot = _num(r[I_TOT])
        otros = sum(_num(r[i]) for i in OTROS)
        base = tot - iva - otros
        if emi == CO:                       # VENTAS (empresa emite)
            k = 'vn' if es_nc else 'vf'
            for f, v in (('n', 1), ('base', base), ('iva', iva), ('otros', otros), ('tot', tot)):
                res[k][f] += v
            tgt = vn if es_nc else vf
            tgt[d]['n'] += 1; tgt[d]['base'] += base; tgt[d]['iva'] += iva
            tgt[d]['otros'] += otros; tgt[d]['tot'] += tot
        elif rec == CO:                     # COMPRAS (empresa recibe)
            k = 'cn' if es_nc else 'cf'
            for f, v in (('n', 1), ('base', base), ('iva', iva), ('otros', otros), ('tot', tot)):
                res[k][f] += v
            p = prov[emi]; p['nom'] = str(r[I_NEMI] or '').strip()
            sgn = -1 if es_nc else 1
            if es_nc:
                p['nc'] += 1
            else:
                p['nf'] += 1
            p['base'] += sgn * base; p['iva'] += sgn * iva
            p['otros'] += sgn * otros; p['tot'] += sgn * tot
    return vf, vn, prov, res


def generar_reporte(token_path, nit_empresa, anio, mes, titulo_empresa, nit_formateado,
                    dia_ini=None, dia_fin=None, salida=None, periodo_texto=None):
    """Genera el .xlsx del reporte recurrente y lo guarda en `salida`. Devuelve la ruta y un dict resumen."""
    vf, vn, prov, res = procesar_token(token_path, nit_empresa, anio, mes, dia_ini, dia_fin)
    if periodo_texto is None:
        meses = ['', 'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO',
                 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']
        rango = f" ({dia_ini} al {dia_fin})" if (dia_ini or dia_fin) else ""
        periodo_texto = f"{meses[mes]} {anio}{rango}"
    out = Workbook()

    # ---------- Hoja 1: Resumen IVA ----------
    w = out.active; w.title = "Resumen IVA"; w.sheet_view.showGridLines = False
    w.merge_cells('A1:F1')
    _S(w['A1'], f"LIQUIDACIÓN DE IVA — {titulo_empresa} · NIT {nit_formateado} · {periodo_texto}",
       b=True, sz=13, col=WHITE, fill=NAVY, al="center"); w.row_dimensions[1].height = 24
    w.merge_cells('A2:F2')
    _S(w['A2'], "Fuente: factura electrónica DIAN · excluye 'Application response' · notas crédito = devoluciones",
       sz=9, col=WHITE, fill=BLUE, al="center")
    for j, h in enumerate(["Concepto", "# Docs", "Base (sin IVA)", "IVA", "Otros impuestos", "Total (con IVA)"], 1):
        _S(w.cell(4, j), h, b=True, col=WHITE, fill=BLUE, al="center", wrap=True)

    def _ln(rr, lbl, dd, b=False, fill=None):
        _S(w.cell(rr, 1), lbl, b=b, fill=fill); _S(w.cell(rr, 2), dd['n'], al="center", b=b, fill=fill)
        _S(w.cell(rr, 3), round(dd['base']), al="right", num=True, b=b, fill=fill)
        _S(w.cell(rr, 4), round(dd['iva']), al="right", num=True, b=b, fill=fill)
        _S(w.cell(rr, 5), round(dd['otros']), al="right", num=True, b=b, fill=fill)
        _S(w.cell(rr, 6), round(dd['tot']), al="right", num=True, b=b, fill=fill)

    _S(w.cell(5, 1), "VENTAS", b=True, fill=GOLD); [_S(w.cell(5, j), "", fill=GOLD) for j in range(2, 7)]
    _ln(6, "Ventas acumuladas (facturas)", res['vf']); _ln(7, "(−) Devoluciones en ventas (N. crédito)", res['vn'])
    for j, c in ((2, 'B'), (3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')):
        _S(w.cell(8, j), f"={c}6-{c}7", al="right" if j > 2 else "center", num=(j > 2), b=True, fill=LBLUE)
    _S(w.cell(8, 1), "(=) VENTAS NETAS", b=True, fill=LBLUE)
    _S(w.cell(10, 1), "COMPRAS", b=True, fill=GOLD); [_S(w.cell(10, j), "", fill=GOLD) for j in range(2, 7)]
    _ln(11, "Compras acumuladas (facturas)", res['cf']); _ln(12, "(−) Devoluciones en compras (N. crédito)", res['cn'])
    for j, c in ((2, 'B'), (3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')):
        _S(w.cell(13, j), f"={c}11-{c}12", al="right" if j > 2 else "center", num=(j > 2), b=True, fill=LBLUE)
    _S(w.cell(13, 1), "(=) COMPRAS NETAS", b=True, fill=LBLUE)
    _S(w.cell(15, 1), "LIQUIDACIÓN DE IVA", b=True, sz=11, fill=GOLD); [_S(w.cell(15, j), "", fill=GOLD) for j in range(2, 7)]
    _S(w.cell(16, 1), "IVA generado (ventas netas)"); w.merge_cells('A16:C16'); _S(w.cell(16, 4), "=D8", al="right", num=True, b=True)
    _S(w.cell(17, 1), "IVA descontable (compras netas)"); w.merge_cells('A17:C17'); _S(w.cell(17, 4), "=D13", al="right", num=True, b=True)
    _S(w.cell(18, 1), "SALDO IVA A PAGAR A LA FECHA", b=True, fill=GREEN, col=GREENF); w.merge_cells('A18:C18')
    _S(w.cell(18, 4), "=D16-D17", al="right", num=True, b=True, fill=GREEN, col=GREENF)
    for col, wd in zip("ABCDEF", [38, 9, 17, 15, 15, 17]):
        w.column_dimensions[col].width = wd

    # ---------- Hoja 2: Ventas diarias ----------
    w2 = out.create_sheet("Ventas diarias"); w2.sheet_view.showGridLines = False
    w2.merge_cells('A1:K1')
    _S(w2['A1'], f"VENTAS DIARIAS POR FECHA — {titulo_empresa} · {periodo_texto}", b=True, sz=12, col=WHITE, fill=NAVY, al="center")
    w2.row_dimensions[1].height = 22
    w2.merge_cells('A2:K2')
    _S(w2['A2'], "Día de la semana al lado de la fecha · venta neta antes de IVA = base ventas − base devoluciones",
       sz=9, col=WHITE, fill=BLUE, al="center")
    hdr = ["Fecha", "Día", "# Facturas", "Base ventas", "IVA ventas", "Total ventas",
           "# N.Créd", "Total devol.", "IVA devol.", "Venta neta (con IVA)", "Venta neta ANTES de IVA"]
    for j, h in enumerate(hdr, 1):
        _S(w2.cell(3, j), h, b=True, col=WHITE, fill=BLUE, al="center", wrap=True)
    w2.row_dimensions[3].height = 30
    rr = 4; first = rr
    for dia in sorted(set(vf) | set(vn)):
        v = vf.get(dia, {'n': 0, 'base': 0, 'iva': 0, 'tot': 0})
        n = vn.get(dia, {'n': 0, 'base': 0, 'iva': 0, 'tot': 0})
        _S(w2.cell(rr, 1), dia.strftime('%Y-%m-%d'), al="center"); _S(w2.cell(rr, 2), DIAS[dia.weekday()], al="center")
        _S(w2.cell(rr, 3), v['n'], al="center", num=True); _S(w2.cell(rr, 4), round(v['base']), al="right", num=True)
        _S(w2.cell(rr, 5), round(v['iva']), al="right", num=True); _S(w2.cell(rr, 6), round(v['tot']), al="right", num=True)
        _S(w2.cell(rr, 7), n['n'], al="center", num=True); _S(w2.cell(rr, 8), round(n['tot']), al="right", num=True)
        _S(w2.cell(rr, 9), round(n['iva']), al="right", num=True)
        _S(w2.cell(rr, 10), f"=F{rr}-H{rr}", al="right", num=True); _S(w2.cell(rr, 11), f"=D{rr}-H{rr}+I{rr}", al="right", num=True)
        rr += 1
    last = rr - 1
    _S(w2.cell(rr, 1), "TOTAL", b=True, fill=LBLUE); _S(w2.cell(rr, 2), "", fill=LBLUE)
    for c in "CDEFGHIJK":
        _S(w2.cell(rr, ord(c) - 64), f"=SUM({c}{first}:{c}{last})", b=True,
           al="center" if c in "CG" else "right", num=True, fill=LBLUE)
    for col, wd in zip("ABCDEFGHIJK", [13, 12, 11, 15, 14, 15, 9, 13, 12, 17, 18]):
        w2.column_dimensions[col].width = wd
    w2.freeze_panes = "A4"

    # ---------- Hoja 3: Compras + IVA por proveedor ----------
    w3 = out.create_sheet("Compras+IVA x proveedor"); w3.sheet_view.showGridLines = False
    w3.merge_cells('A1:G1')
    _S(w3['A1'], f"COMPRAS + IVA POR PROVEEDOR — {titulo_empresa} · {periodo_texto}", b=True, sz=12, col=WHITE, fill=NAVY, al="center")
    w3.row_dimensions[1].height = 22
    w3.merge_cells('A2:G2')
    _S(w3['A2'], "Neto de notas crédito · ordenado por compra + IVA", sz=9, col=WHITE, fill=BLUE, al="center")
    for j, h in enumerate(["NIT", "Proveedor", "# Fact", "# N.Créd", "Base compra", "IVA", "Compra + IVA"], 1):
        _S(w3.cell(3, j), h, b=True, col=WHITE, fill=BLUE, al="center", wrap=True)
    rr = 4; first = rr
    for nit, p in sorted(prov.items(), key=lambda kv: -(kv[1]['base'] + kv[1]['iva'] + kv[1]['otros'])):
        _S(w3.cell(rr, 1), nit, al="center"); _S(w3.cell(rr, 2), p['nom'])
        _S(w3.cell(rr, 3), p['nf'], al="center"); _S(w3.cell(rr, 4), p['nc'], al="center")
        _S(w3.cell(rr, 5), round(p['base']), al="right", num=True); _S(w3.cell(rr, 6), round(p['iva']), al="right", num=True)
        _S(w3.cell(rr, 7), f"=E{rr}+F{rr}", al="right", num=True)
        rr += 1
    last = rr - 1
    _S(w3.cell(rr, 2), "TOTAL", b=True, fill=LBLUE, al="right")
    for c in "CDEFG":
        _S(w3.cell(rr, ord(c) - 64), f"=SUM({c}{first}:{c}{last})", b=True,
           al="center" if c in "CD" else "right", num=True, fill=LBLUE)
    for col, wd in zip("ABCDEFG", [12, 34, 8, 9, 16, 15, 16]):
        w3.column_dimensions[col].width = wd
    w3.freeze_panes = "A4"

    if salida is None:
        salida = f"Reporte_{nit_empresa}_{anio}-{mes:02d}.xlsx"
    out.save(salida)
    iva_gen = res['vf']['iva'] - res['vn']['iva']
    iva_desc = res['cf']['iva'] - res['cn']['iva']
    resumen = {
        'ventas_netas_base': res['vf']['base'] - res['vn']['base'],
        'iva_generado': iva_gen, 'iva_descontable': iva_desc,
        'iva_a_pagar': iva_gen - iva_desc, 'dias': len(vf), 'proveedores': len(prov),
    }
    return salida, resumen


def generar_reporte_bytes(token_path, nit_empresa, anio, mes, titulo_empresa, nit_formateado,
                          dia_ini=None, dia_fin=None, periodo_texto=None):
    """Igual que generar_reporte pero devuelve (bytes_xlsx, resumen) — ideal para Streamlit."""
    import tempfile, os
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    _, resumen = generar_reporte(token_path, nit_empresa, anio, mes, titulo_empresa,
                                 nit_formateado, dia_ini, dia_fin, salida=tmp.name,
                                 periodo_texto=periodo_texto)
    with open(tmp.name, "rb") as f:
        data = f.read()
    os.unlink(tmp.name)
    return data, resumen


# Catálogo de empresas conocidas (opcional; para autocompletar título/NIT formateado)
EMPRESAS = {
    "900451388": {"titulo": "SILLA TRES S.A.S", "nit_fmt": "900.451.388"},
}
