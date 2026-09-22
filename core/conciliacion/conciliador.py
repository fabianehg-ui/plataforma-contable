# -*- coding: utf-8 -*-
"""
core/conciliacion/conciliador.py

Conciliación bancaria config-driven para INTEGRAL.

Cruza tres fuentes por cuenta bancaria:
  1) AUXILIAR general del mes (libros): filtrado a la cuenta puente del banco
     (p.ej. 11-10-05-99). De aquí salen el SALDO EN LIBROS (saldo anterior),
     las CONSIGNACIONES (total débitos) y los PAGOS (total créditos).
  2) REPORTE DEL BANCO: la hoja del banco con los movimientos y la columna
     'N COMPROBANTE' (el número que la empresa coloca). Se cruza N COMPROBANTE
     con el 'Documento' del auxiliar. Los GTO FINANCIERO se desglosan en gasto
     bancario / comisiones / IVA / GMF.
  3) DATÁFONO (Credibanco, opcional): comisión y retenciones (retefuente,
     reteIVA, reteICA) por centro de costo, que también son NOTAS DÉBITO.

Fórmula (igual a la plantilla de la empresa):
    Saldo en libros
    + Consignaciones (débitos auxiliar)
    − Pagos (créditos auxiliar)
    − Notas débito (gasto banc + comis + comis datáfono + IVA + GMF +
                    reteIVA + retefuente + reteICA)
    = Saldo en libros a fecha de cierre
    − Consignaciones en tránsito (partida conciliatoria)
    = Saldo según extracto del banco

Partidas conciliatorias:
    - En tránsito / no en banco : documentos en libros que no están en el banco.
    - Pagos que no salieron     : ídem, del lado de pagos.
    - No están en contabilidad  : documentos en el banco que no están en libros.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict


def _num(x) -> float:
    try:
        return float(str(x).replace(",", "").replace("$", "").strip() or 0)
    except (ValueError, TypeError):
        return 0.0


def _nk(x) -> str:
    """Llave numérica normalizada de un documento/comprobante (sin ceros a la izq.)."""
    return re.sub(r"\D", "", str(x or "")).lstrip("0")


def _abrir(fuente, hoja=None):
    import openpyxl
    if isinstance(fuente, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(fuente), data_only=True, read_only=True)
    elif hasattr(fuente, "read"):
        wb = openpyxl.load_workbook(io.BytesIO(fuente.read()), data_only=True, read_only=True)
    else:
        wb = openpyxl.load_workbook(fuente, data_only=True, read_only=True)
    if hoja is None:
        return wb.active, wb
    for ws in wb.worksheets:
        if ws.title.strip().upper() == str(hoja).strip().upper():
            return ws, wb
    return wb.active, wb


# ---- reglas de desglose de GTO FINANCIERO (editable) ----
GASTO_REGLAS_DEF = [
    (r"4X1000|\bGMF\b|IMPTO GOBIERNO|GRAVAMEN", "GMF"),
    (r"\bIVA\b", "IVA"),
    (r"COMIS", "COMISIONES"),
    # el resto -> GASTO BANCARIO
]


def leer_auxiliar(fuente, cuenta_prefijo: str):
    """Del AUXILIAR: saldo anterior, consignaciones (deb), pagos (cred) y el
    detalle por documento de la cuenta que empieza por `cuenta_prefijo`.
    Ignora la fila de SUBTOTAL de la cuenta (la que no trae Nro Registro)."""
    ws, _ = _abrir(fuente)
    rows = list(ws.iter_rows(values_only=True))
    # localizar columnas por encabezado (fila con 'Cuenta' y 'Documento')
    hdr_i = None
    for i, r in enumerate(rows[:8]):
        vals = [str(c or "").strip().lower() for c in r]
        if "cuenta" in vals and "documento" in vals:
            hdr_i = i
            break
    if hdr_i is None:
        hdr_i = 2
    H = {str(c or "").strip().lower(): j for j, c in enumerate(rows[hdr_i])}
    cCta = H.get("cuenta", 0)
    cReg = H.get("nro registro", 5)
    cDoc = H.get("documento", 8)
    cFec = H.get("fecha", 7)
    cDet = H.get("detalle", 10)
    cSA = H.get("saldo anterior", 12)
    cDeb = H.get("débitos", H.get("debitos", 13))
    cCred = H.get("créditos", H.get("creditos", 14))

    pref = cuenta_prefijo.strip()
    saldo_ant = None
    consig = pagos = 0.0
    detalle = []
    for r in rows[hdr_i + 1:]:
        cta = str(r[cCta] or "").strip()
        if not cta.startswith(pref):
            continue
        if saldo_ant is None:
            saldo_ant = _num(r[cSA])
        if not str(r[cReg] or "").strip():
            continue  # fila subtotal de la cuenta
        deb, cred = _num(r[cDeb]), _num(r[cCred])
        if deb == 0 and cred == 0:
            continue
        consig += deb
        pagos += cred
        detalle.append({"doc": _nk(r[cDoc]), "fecha": str(r[cFec] or ""),
                        "detalle": str(r[cDet] or "")[:60], "deb": deb, "cred": cred})
    return {"saldo_ant": saldo_ant or 0.0, "consignaciones": round(consig, 2),
            "pagos": round(pagos, 2), "detalle": detalle}


def leer_banco(fuente, hoja: str, gasto_reglas=None):
    """De la HOJA del banco: movimientos con N COMPROBANTE, débito, crédito y
    concepto; desglose de GTO FINANCIERO en gasto/comisiones/IVA/GMF."""
    reglas = gasto_reglas or GASTO_REGLAS_DEF
    ws, _ = _abrir(fuente, hoja)
    rows = list(ws.iter_rows(values_only=True))
    # encabezado: primera fila con 'N COMPROBANTE' o 'DEBITO'
    hdr_i = 0
    for i, r in enumerate(rows[:6]):
        up = [str(c or "").strip().upper() for c in r]
        if "N COMPROBANTE" in up or "DEBITO" in up or "DÉBITO" in up:
            hdr_i = i
            break
    H = {str(c or "").strip().upper(): j for j, c in enumerate(rows[hdr_i])}
    cNC = H.get("N COMPROBANTE")
    cDeb = H.get("DEBITO", H.get("DÉBITO", H.get("DEBITOS", H.get("DÉBITOS"))))
    cCred = H.get("CREDITO", H.get("CRÉDITO", H.get("CREDITOS", H.get("CRÉDITOS"))))
    cCon = H.get("CONCEPTO")
    cDet = H.get("DETALLE", H.get("TRANSACCIÓN", H.get("DESCRIPCIÓN")))
    cFec = H.get("FECHA")

    movs = []
    g = defaultdict(float)
    for r in rows[hdr_i + 1:]:
        deb = _num(r[cDeb]) if cDeb is not None else 0.0
        cred = _num(r[cCred]) if cCred is not None else 0.0
        if deb == 0 and cred == 0:
            continue
        det = str(r[cDet] or "") if cDet is not None else ""
        con = str(r[cCon] or "").strip() if cCon is not None else ""
        movs.append({"nc": _nk(r[cNC]) if cNC is not None else "",
                     "fecha": str(r[cFec] or "") if cFec is not None else "",
                     "detalle": det[:60], "deb": deb, "cred": cred, "concepto": con})
        if con.upper() == "GTO FINANCIERO":
            v = deb + cred
            cat = "GASTO BANCARIO"
            for pat, nombre in reglas:
                if re.search(pat, det, re.I):
                    cat = nombre
                    break
            g[cat] += v
    return {"movimientos": movs, "gastos": {k: round(v, 2) for k, v in g.items()}}


def _es_encab_datafono(fila):
    """¿Esta fila es el encabezado del resumen del datáfono? Devuelve el mapeo
    de columnas {clave: indice} si lo es, o None."""
    up = [str(c or "").strip().upper() for c in fila]
    joined = " | ".join(up)
    if not (("COMIS" in joined) and ("CENTRO" in joined or "COSTO" in joined)):
        return None
    mp = {}
    for j, v in enumerate(up):
        if ("CENTRO" in v or "COSTO" in v) and "cc" not in mp:
            mp["cc"] = j
        elif "OASIS" in v or "ESTABLEC" in v or "NOMBRE" in v or "PUNTO" in v:
            mp.setdefault("oasis", j)
        elif "COMIS" in v:
            mp.setdefault("comision", j)
        elif "RETEFUENTE" in v or ("RTE" in v and "FUENTE" in v) or "RETE FUENTE" in v \
                or ("RETE" in v and "FUENTE" in v):
            mp.setdefault("retefuente", j)
        elif "IVA" in v and ("RETE" in v or "RTE" in v):
            mp.setdefault("reteiva", j)
        elif "ICA" in v and ("RETE" in v or "RTE" in v):
            mp.setdefault("reteica", j)
    return mp if "comision" in mp else None


def leer_datafono(fuente, hoja="RESUMEN MENSUAL"):
    """Del macro Credibanco: comisión y retenciones por centro de costo.

    Robusto: busca la hoja del resumen y detecta las columnas por su encabezado
    (CENTRO DE COSTO / GASTO COMISION / RETEFUENTE / RETE IVA / RTE ICA), sin
    depender de que la hoja se llame exactamente 'RESUMEN MENSUAL' ni de que las
    columnas estén en posiciones fijas. Ignora la fila de totales."""
    vacio = {"por_cc": [], "comision": 0.0, "retefuente": 0.0,
             "reteiva": 0.0, "reteica": 0.0, "n_filas": 0, "hoja": None}
    if fuente is None:
        return vacio

    import openpyxl
    if isinstance(fuente, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(fuente), data_only=True, read_only=True)
    elif hasattr(fuente, "read"):
        wb = openpyxl.load_workbook(io.BytesIO(fuente.read()), data_only=True, read_only=True)
    else:
        wb = openpyxl.load_workbook(fuente, data_only=True, read_only=True)

    # candidatas: primero la hoja pedida, luego el resto
    orden = ([ws for ws in wb.worksheets if ws.title.strip().upper() == str(hoja).strip().upper()]
             + [ws for ws in wb.worksheets if ws.title.strip().upper() != str(hoja).strip().upper()])

    for ws in orden:
        rows = list(ws.iter_rows(values_only=True))
        mp = hdr_i = None
        for i, r in enumerate(rows[:15]):
            m = _es_encab_datafono(r)
            if m:
                mp, hdr_i = m, i
                break
        if not mp:
            continue

        cCC, cCom = mp.get("cc", 0), mp["comision"]
        cOa = mp.get("oasis", cCC + 1)
        cRF, cRI, cRA = mp.get("retefuente"), mp.get("reteiva"), mp.get("reteica")
        por_cc = []
        com = ret = riva = rica = 0.0
        for r in rows[hdr_i + 1:]:
            cc = str(r[cCC] or "").strip() if cCC < len(r) else ""
            if not cc or cc.upper().startswith(("TOTAL", "SUMA")):
                continue
            # son NOTAS DÉBITO (cargos): se toman en magnitud, por si el export
            # de Credibanco los trae con signo negativo (p.ej. JIPER).
            c = abs(_num(r[cCom])) if cCom < len(r) else 0.0
            rf = abs(_num(r[cRF])) if cRF is not None and cRF < len(r) else 0.0
            ri = abs(_num(r[cRI])) if cRI is not None and cRI < len(r) else 0.0
            ria = abs(_num(r[cRA])) if cRA is not None and cRA < len(r) else 0.0
            if c == 0 and rf == 0 and ri == 0 and ria == 0:
                continue
            oa = str(r[cOa] or "") if cOa < len(r) else ""
            por_cc.append({"cc": cc, "oasis": oa, "comision": c,
                           "retefuente": rf, "reteiva": ri, "reteica": ria})
            com += c; ret += rf; riva += ri; rica += ria
        if por_cc:
            return {"por_cc": por_cc, "comision": round(com, 2),
                    "retefuente": round(ret, 2), "reteiva": round(riva, 2),
                    "reteica": round(rica, 2), "n_filas": len(por_cc), "hoja": ws.title}
    return vacio


def conciliar(aux, banco, datafono, saldo_banco: float,
              transito_real=None, tolerancia: float = 10000.0):
    """Arma la conciliación y las partidas. Recibe los dict de leer_* y el saldo
    del extracto del banco.

    - Si `transito_real` es None: las consignaciones en tránsito se calculan como
      el plug (saldo a cierre − saldo del banco), así siempre cuadra.
    - Si se entrega `transito_real` (valor de las partidas reales en puente): el
      residuo = saldo a cierre − tránsito real − saldo del banco se carga como
      AJUSTE AL PESO a gasto bancario SIEMPRE QUE |residuo| ≤ `tolerancia`
      (por defecto ±10.000), para que cuadre exacto. Si supera la tolerancia se
      deja el residuo visible y `cuadra=False` para que lo revises.
    """
    g = banco["gastos"]
    gasto_banc = g.get("GASTO BANCARIO", 0.0)
    comis = g.get("COMISIONES", 0.0)
    iva = g.get("IVA", 0.0)
    gmf = g.get("GMF", 0.0)
    com_d = datafono["comision"]
    ret_d = datafono["retefuente"]
    riva_d = datafono["reteiva"]
    rica_d = datafono["reteica"]

    ajuste = 0.0
    if transito_real is None:
        notas = round(gasto_banc + comis + com_d + iva + gmf + riva_d + ret_d + rica_d, 2)
        cierre = round(aux["saldo_ant"] + aux["consignaciones"] - aux["pagos"] - notas, 2)
        transito = round(cierre - saldo_banco, 2)
    else:
        transito = round(float(transito_real), 2)
        notas0 = round(gasto_banc + comis + com_d + iva + gmf + riva_d + ret_d + rica_d, 2)
        cierre0 = round(aux["saldo_ant"] + aux["consignaciones"] - aux["pagos"] - notas0, 2)
        residuo = round(cierre0 - transito - saldo_banco, 2)   # lo que sobra/falta
        if abs(residuo) <= tolerancia:
            ajuste = residuo                 # se carga a gasto bancario
            gasto_banc = round(gasto_banc + ajuste, 2)
        notas = round(gasto_banc + comis + com_d + iva + gmf + riva_d + ret_d + rica_d, 2)
        cierre = round(aux["saldo_ant"] + aux["consignaciones"] - aux["pagos"] - notas, 2)

    # cruce por documento
    lby = defaultdict(list)
    bby = defaultdict(list)
    for x in aux["detalle"]:
        lby[x["doc"]].append(x)
    for x in banco["movimientos"]:
        bby[x["nc"]].append(x)
    dl = set(lby) - {""}
    db = set(bby) - {""}
    cruzan = sorted(dl & db)
    solo_libros = sorted(dl - db)   # en libros, no en banco
    solo_banco = sorted(db - dl)    # en banco, no en libros (no en contabilidad)

    return {
        "saldo_ant": round(aux["saldo_ant"], 2),
        "consignaciones": aux["consignaciones"],
        "pagos": aux["pagos"],
        "notas_debito": notas,
        "gasto_bancario": gasto_banc, "comisiones": comis,
        "comision_datafono": com_d, "iva": iva, "gmf": gmf,
        "reteiva": riva_d, "retefuente": ret_d, "reteica": rica_d,
        "saldo_cierre": cierre,
        "consignaciones_transito": transito,
        "saldo_banco": round(saldo_banco, 2),
        "ajuste_al_peso": round(ajuste, 2),
        "cuadra": abs((cierre - transito) - saldo_banco) < 0.01,
        "n_cruzan": len(cruzan),
        "solo_libros": [x for d in solo_libros for x in lby[d]],
        "solo_banco": [x for d in solo_banco for x in bby[d]],
        "datafono_cc": datafono["por_cc"],
    }


# ===========================================================================
# Configuración por empresa (Supabase) — tabla conciliacion_bancos
# ===========================================================================
def cargar_bancos(sb, empresa_id) -> list:
    try:
        r = (sb.table("conciliacion_bancos").select("*")
             .eq("empresa_id", empresa_id).order("orden").execute())
        return [dict(x) for x in (r.data or [])]
    except Exception:  # noqa: BLE001
        return []


def guardar_bancos(sb, empresa_id, filas) -> int:
    filas = [f for f in filas if str(f.get("nombre") or "").strip()
             and str(f.get("cuenta_auxiliar") or "").strip()]
    nombres = [str(f["nombre"]).strip() for f in filas]
    for b in cargar_bancos(sb, empresa_id):
        if b["nombre"] not in nombres:
            sb.table("conciliacion_bancos").delete().eq("empresa_id", empresa_id).eq("nombre", b["nombre"]).execute()
    payload = [{"empresa_id": empresa_id, "nombre": str(f["nombre"]).strip(),
                "cuenta_auxiliar": str(f.get("cuenta_auxiliar") or "").strip(),
                "hoja_reporte": str(f.get("hoja_reporte") or "").strip(), "orden": i}
               for i, f in enumerate(filas)]
    if payload:
        sb.table("conciliacion_bancos").upsert(payload, on_conflict="empresa_id,nombre").execute()
    return len(payload)


BANCOS_LOLITA = [
    {"nombre": "BANCOLOMBIA CTE 4451",   "cuenta_auxiliar": "11-10-05-99", "hoja_reporte": "CTA- PUENTE 4451"},
    {"nombre": "OCCIDENTE 9426",         "cuenta_auxiliar": "11-10-05-05", "hoja_reporte": "OCCIDENTE 9426"},
    {"nombre": "DAVIVIENDA 7872",        "cuenta_auxiliar": "11-20-05-13", "hoja_reporte": "Davivienda cta 7872"},
    {"nombre": "BOGOTA 3199",            "cuenta_auxiliar": "11-10-05-20", "hoja_reporte": "BOGOTA 3199"},
    {"nombre": "BANCOLOMBIA AHORROS",    "cuenta_auxiliar": "11-10-05-06", "hoja_reporte": "BANCOLOMBA AHORROS"},
]


def sembrar_lolita(sb, empresa_id) -> int:
    return guardar_bancos(sb, empresa_id, BANCOS_LOLITA)
